"""
Professional multi-sheet Excel export for IFRS 2 SBP Valuation App.
Uses openpyxl with corporate formatting.
"""

import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side)
from openpyxl.utils import get_column_letter

# ── Style constants ──────────────────────────────────────────────────────────
NAVY    = "0B1F38"
BLUE    = "2563EB"
LTBLUE  = "EFF6FF"
WHITE   = "FFFFFF"
GRAY    = "F8FAFC"
BORDER  = "E2E8F0"

H_FILL  = PatternFill("solid", fgColor=NAVY)
SH_FILL = PatternFill("solid", fgColor=BLUE)
A_FILL  = PatternFill("solid", fgColor=LTBLUE)
G_FILL  = PatternFill("solid", fgColor="F1F5F9")
H_FONT  = Font(name="Calibri", bold=True, color=WHITE, size=11)
SH_FONT = Font(name="Calibri", bold=True, color=WHITE, size=10)
BD_FONT = Font(name="Calibri", size=10)
BLD     = Font(name="Calibri", bold=True, size=10)
THIN    = Border(
    left=Side(style="thin", color=BORDER),
    right=Side(style="thin", color=BORDER),
    top=Side(style="thin", color=BORDER),
    bottom=Side(style="thin", color=BORDER),
)

FMT_NUM = '#,##0.00'
FMT_PCT = '0.00%'
FMT_INT = '#,##0'
FMT_NUM4 = '#,##0.0000'


def _cw(ws, col, w):
    ws.column_dimensions[get_column_letter(col)].width = w


def _hr(ws, row, vals, fill=None, font=None, height=18, cols=None):
    fill = fill or H_FILL
    font = font or H_FONT
    ws.row_dimensions[row].height = height
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row, c, v)
        cell.fill = fill; cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN


def _dr(ws, row, vals, fmts=None, bold=False):
    fmts = fmts or [None] * len(vals)
    alt = row % 2 == 0
    for c, (v, f) in enumerate(zip(vals, fmts), 1):
        cell = ws.cell(row, c, v)
        cell.font = BLD if bold else BD_FONT
        cell.border = THIN
        if f: cell.number_format = f
        if alt: cell.fill = G_FILL
        cell.alignment = Alignment(vertical="center")


def _title(ws, row, text, cols=4):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    c = ws.cell(row, 1, text)
    c.fill = SH_FILL; c.font = SH_FONT
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.border = THIN
    ws.row_dimensions[row].height = 16


# ════════════════════════════════════════════════════════════════════════════

def build_excel(
    inputs: dict,
    bs_results: dict | None,
    sigma: float | None,
    first_year: dict | None,
    vesting_sched: list,
    accounting_entries: list,
    subsequent_result: dict | None,
    log_returns_series=None,
    prices_df=None,
    vol_info: dict | None = None,
    summary_kpis: dict | None = None,
) -> bytes:

    wb = Workbook()
    wb.remove(wb.active)

    # ── 1. Executive Summary ─────────────────────────────────────────────────
    ws = wb.create_sheet("Executive Summary")
    _cw(ws, 1, 38); _cw(ws, 2, 22); _cw(ws, 3, 22); _cw(ws, 4, 40)
    _hr(ws, 1, ["IFRS 2 Share-Based Payment – Executive Summary", "", "", ""], cols=4)
    ws.merge_cells("A1:D1")

    row = 3
    _title(ws, row, "KEY PERFORMANCE INDICATORS", 4); row += 1
    _hr(ws, row, ["Metric", "Value", "", "Note"], fill=SH_FILL, font=SH_FONT); row += 1
    kpis = summary_kpis or {}
    for k, v in kpis.items():
        _dr(ws, row, [k, v, "", ""])
        row += 1

    # ── 2. Inputs ────────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Inputs")
    _cw(ws2, 1, 38); _cw(ws2, 2, 20); _cw(ws2, 3, 45)
    _hr(ws2, 1, ["Parameter", "Value", "Description"])
    row = 2
    _title(ws2, row, "PROGRAM SETUP", 3); row += 1
    setup_rows = [
        ("Program Type", inputs.get("program_type", ""), "Equity-settled / Cash-settled / Employee choice"),
        ("Instrument Type", inputs.get("instrument_type", ""), "Stock option / Warrant / RSU"),
        ("Grant Date", str(inputs.get("grant_date", "")), "Date options are granted"),
        ("Reporting Date", str(inputs.get("reporting_date", "")), "Current reporting date"),
        ("Vesting Period (years)", inputs.get("vesting_period", ""), "Service/vesting period"),
        ("Total Program Period (years)", inputs.get("total_period", ""), "Expected life"),
        ("Number of Employees", inputs.get("n_employees", ""), ""),
        ("Options per Employee", inputs.get("options_per_employee", ""), ""),
        ("Total Instruments Granted", inputs.get("n_total", ""), ""),
        ("Forfeiture Rate", inputs.get("forfeiture_rate", ""), ""),
        ("Expected Vesting %", inputs.get("expected_vesting_pct", ""), ""),
        ("Months in First Year", inputs.get("months_in_first_year", 12), "For proration"),
    ]
    for p, v, d in setup_rows:
        fmt = FMT_PCT if "Rate" in p or "%" in p else None
        _dr(ws2, row, [p, v, d], fmts=[None, fmt, None]); row += 1

    row += 1
    _title(ws2, row, "BLACK-SCHOLES INPUTS", 3); row += 1
    bs_inp = [
        ("Share Price S₀", inputs.get("S0", ""), "Spot price at grant / reporting date"),
        ("Strike Price K", inputs.get("K", ""), "Exercise price"),
        ("Risk-free Rate r", inputs.get("r", ""), "Annual, continuous"),
        ("Dividend Yield q", inputs.get("q", ""), "Annual, continuous"),
        ("Volatility σ", sigma, "Annualized"),
        ("Expected Life T (years)", inputs.get("T", ""), ""),
    ]
    for p, v, d in bs_inp:
        fmt = FMT_PCT if p in ("Risk-free Rate r", "Dividend Yield q", "Volatility σ") else FMT_NUM if isinstance(v, float) else None
        _dr(ws2, row, [p, v, d], fmts=[None, fmt, None]); row += 1

    # ── 3. Historical Prices ─────────────────────────────────────────────────
    ws3 = wb.create_sheet("Historical Prices")
    if prices_df is not None and not prices_df.empty:
        cols = list(prices_df.columns)
        _hr(ws3, 1, cols + [""])
        for c, col in enumerate(cols, 1):
            _cw(ws3, c, 16)
        for i, row_data in enumerate(prices_df.itertuples(index=False), 2):
            _dr(ws3, i, list(row_data))
    else:
        ws3.cell(1, 1, "No historical prices uploaded.").font = BD_FONT

    # ── 4. Log Returns ───────────────────────────────────────────────────────
    ws4 = wb.create_sheet("Log Returns")
    _cw(ws4, 1, 15); _cw(ws4, 2, 20)
    _hr(ws4, 1, ["Index", "Log Return ln(Pt/Pt-1)"])
    if log_returns_series is not None and len(log_returns_series) > 0:
        for i, (idx, val) in enumerate(zip(log_returns_series.index, log_returns_series.values), 2):
            _dr(ws4, i, [str(idx), float(val)], fmts=[None, FMT_NUM4])
    else:
        ws4.cell(2, 1, "No log returns calculated.").font = BD_FONT

    # ── 5. Volatility ────────────────────────────────────────────────────────
    ws5 = wb.create_sheet("Volatility Calculation")
    _cw(ws5, 1, 32); _cw(ws5, 2, 20); _cw(ws5, 3, 35)
    _hr(ws5, 1, ["Item", "Value", "Formula"])
    row = 2
    if vol_info:
        vol_rows = [
            ("Frequency", vol_info.get("frequency", ""), "daily / weekly / monthly"),
            ("Std Dev of Log Returns", vol_info.get("std", ""), "Sample std dev of ln(Pt/Pt-1)"),
            ("Annualization Factor", vol_info.get("factor", ""), "√252 / √52 / √12"),
            ("Annualized Volatility σ", vol_info.get("ann_vol", ""), "Std × Factor"),
        ]
        for p, v, f in vol_rows:
            fmt = FMT_PCT if "Volatility" in p else FMT_NUM4 if isinstance(v, float) else None
            _dr(ws5, row, [p, v, f], fmts=[None, fmt, None]); row += 1
    else:
        ws5.cell(2, 1, f"Manual input: {sigma:.4f}" if sigma else "N/A").font = BD_FONT

    # ── 6. Black-Scholes Valuation ───────────────────────────────────────────
    ws6 = wb.create_sheet("Black-Scholes Valuation")
    _cw(ws6, 1, 32); _cw(ws6, 2, 22); _cw(ws6, 3, 45)
    _hr(ws6, 1, ["Parameter", "Value", "Formula / Notes"])
    row = 2
    if bs_results:
        bsr = [
            ("d₁", bs_results.get("d1"), "[ ln(S₀/K) + (r-q+σ²/2)·T ] / (σ·√T)"),
            ("d₂", bs_results.get("d2"), "d₁ - σ·√T"),
            ("N(d₁)", bs_results.get("Nd1"), "Standard normal CDF of d₁"),
            ("N(d₂)", bs_results.get("Nd2"), "Standard normal CDF of d₂"),
            ("N(-d₁)", bs_results.get("N_neg_d1"), ""),
            ("N(-d₂)", bs_results.get("N_neg_d2"), ""),
            ("Call Option Price", bs_results.get("call"), "S₀·e^(-qT)·N(d₁) - K·e^(-rT)·N(d₂)"),
            ("Put Option Price", bs_results.get("put"), "K·e^(-rT)·N(-d₂) - S₀·e^(-qT)·N(-d₁)"),
            ("Delta (Call)", bs_results.get("delta_call"), "∂C/∂S"),
            ("Delta (Put)", bs_results.get("delta_put"), "∂P/∂S"),
            ("Gamma", bs_results.get("gamma"), "∂²C/∂S²"),
            ("Vega (per 1%σ)", bs_results.get("vega"), "∂C/∂σ per 1%"),
            ("Theta (Call, per day)", bs_results.get("theta_call"), "Time decay per calendar day"),
        ]
        for p, v, n in bsr:
            fmt = FMT_NUM4 if isinstance(v, float) else None
            _dr(ws6, row, [p, v, n], fmts=[None, fmt, None]); row += 1

    # ── 7. Vesting Assumptions ───────────────────────────────────────────────
    ws7 = wb.create_sheet("Vesting Assumptions")
    _cw(ws7, 1, 30); _cw(ws7, 2, 18); _cw(ws7, 3, 18)
    _hr(ws7, 1, ["Year", "Period Expense", "Cumulative Expense"])
    for i, r in enumerate(vesting_sched, 2):
        _dr(ws7, i, [r["Year"], r["Expense"], r["Cumulative"]], fmts=[FMT_INT, FMT_NUM, FMT_NUM])

    # ── 8. First-Year Accounting ─────────────────────────────────────────────
    ws8 = wb.create_sheet("First-Year Accounting")
    _cw(ws8, 1, 38); _cw(ws8, 2, 20); _cw(ws8, 3, 40)
    _hr(ws8, 1, ["Item", "Amount", "Formula"])
    row = 2
    if first_year:
        fy_rows = [
            ("Grant-Date Fair Value per Option", first_year.get("grant_fv"), "Black-Scholes call"),
            ("Expected Vested Instruments", first_year.get("expected_vested"), "Granted × (1 - forfeiture)"),
            ("Total Award Fair Value", first_year.get("total_fv"), "FV × Expected Vested"),
            ("Annual Expense", first_year.get("annual_expense"), "Total FV / Vesting Period"),
            ("Months Active in Year 1", first_year.get("months"), "Proration input"),
            ("First-Year Expense", first_year.get("first_year_expense"), "Annual × Months/12"),
        ]
        for p, v, f in fy_rows:
            fmt = FMT_INT if p == "Expected Vested Instruments" or p == "Months Active in Year 1" else FMT_NUM if isinstance(v, (int, float)) else None
            _dr(ws8, row, [p, v, f], fmts=[None, fmt, None]); row += 1

    # ── 9. Accounting Entries ────────────────────────────────────────────────
    ws9 = wb.create_sheet("Accounting Entries")
    _cw(ws9, 1, 42); _cw(ws9, 2, 18); _cw(ws9, 3, 18); _cw(ws9, 4, 50)
    _hr(ws9, 1, ["Account", "Debit", "Credit", "IFRS 2 Reference"])
    for i, e in enumerate(accounting_entries, 2):
        dr_val = e.debit if e.debit else None
        cr_val = e.credit if e.credit else None
        _dr(ws9, i, [e.account, dr_val, cr_val, e.explanation],
            fmts=[None, FMT_NUM, FMT_NUM, None])

    # ── 10. Subsequent Period ────────────────────────────────────────────────
    ws10 = wb.create_sheet("Subsequent Period")
    _cw(ws10, 1, 40); _cw(ws10, 2, 22); _cw(ws10, 3, 40)
    _hr(ws10, 1, ["Item", "Value", "Notes"])
    row = 2
    if subsequent_result:
        for k, v in subsequent_result.items():
            if k == "entries" or k == "bs_results":
                continue
            fmt = FMT_NUM if isinstance(v, float) else FMT_PCT if "pct" in k or "rate" in k else None
            _dr(ws10, row, [k.replace("_", " ").title(), v, ""], fmts=[None, fmt, None])
            row += 1
        # sub entries
        if "entries" in subsequent_result:
            row += 1
            _title(ws10, row, "JOURNAL ENTRIES", 3); row += 1
            _hr(ws10, row, ["Account", "Debit", "Credit"], fill=SH_FILL, font=SH_FONT); row += 1
            for e in subsequent_result["entries"]:
                _dr(ws10, row, [e.account,
                                e.debit if e.debit else None,
                                e.credit if e.credit else None],
                    fmts=[None, FMT_NUM, FMT_NUM])
                row += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
